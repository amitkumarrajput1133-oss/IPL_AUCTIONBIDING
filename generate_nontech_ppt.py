import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    BG_DARK = RGBColor(11, 19, 43)        # #0B132B Deep Navy
    PANEL_BG = RGBColor(23, 37, 84)       # #172554 Slate Blue Panel
    PANEL_BORDER = RGBColor(51, 65, 85)   # #334155 Border
    CYAN = RGBColor(56, 189, 248)         # #38BDF8 Accent Cyan
    GOLD = RGBColor(245, 158, 11)         # #F59E0B Accent Gold
    GREEN = RGBColor(16, 185, 129)        # #10B981 Live Green
    RED = RGBColor(239, 68, 68)           # #EF4444 Alert Red
    TEXT_LIGHT = RGBColor(248, 250, 252)  # #F8FAFC Main Text
    TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8 Secondary Text

    FONT_TITLE = "Segoe UI"
    FONT_BODY = "Segoe UI"

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text, speaker_name, slide_num):
        # Category / Kicker
        kicker_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.5), Inches(0.4))
        tf_k = kicker_box.text_frame
        tf_k.word_wrap = True
        tf_k.margin_left = tf_k.margin_top = tf_k.margin_right = tf_k.margin_bottom = 0
        p_k = tf_k.paragraphs[0]
        p_k.text = category_text.upper()
        p_k.font.size = Pt(11)
        p_k.font.bold = True
        p_k.font.color.rgb = CYAN
        p_k.font.name = FONT_TITLE

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(8.5), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_LIGHT
        p_t.font.name = FONT_TITLE

        # Speaker Badge (Top Right)
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.6), Inches(0.55), Inches(2.9), Inches(0.7))
        badge.fill.solid()
        badge.fill.fore_color.rgb = PANEL_BG
        badge.line.color.rgb = GOLD
        badge.line.width = Pt(1.5)
        tf_b = badge.text_frame
        tf_b.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_b = tf_b.paragraphs[0]
        p_b.text = f"🎤 SPEAKER: {speaker_name}\nSlide {slide_num} of 8"
        p_b.alignment = PP_ALIGN.CENTER
        p_b.font.size = Pt(10.5)
        p_b.font.bold = True
        p_b.font.color.rgb = GOLD
        p_b.font.name = FONT_TITLE

    def create_card(slide, left, top, width, height, title, items, badge_text=None, border_color=PANEL_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = PANEL_BG
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)
        tf.margin_bottom = Inches(0.2)

        p0 = tf.paragraphs[0]
        if badge_text:
            p0.text = f"[{badge_text}] {title}"
        else:
            p0.text = title
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = CYAN
        p0.font.name = FONT_TITLE
        p0.space_after = Pt(12)

        for item in items:
            p = tf.add_paragraph()
            p.text = item
            p.font.size = Pt(12.5)
            p.font.color.rgb = TEXT_LIGHT
            p.font.name = FONT_BODY
            p.space_after = Pt(8)

        return card

    # SLIDE 1: TITLE SLIDE
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    kicker = s1.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(0.5))
    tf1 = kicker.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "SYNTAX TERROR  •  B.TECH CSE (IOT & BLOCKCHAIN)  •  5TH SEMESTER 2025-26"
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = GOLD
    p1.font.name = FONT_TITLE

    title1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(11.3), Inches(1.6))
    tf_t1 = title1.text_frame
    tf_t1.word_wrap = True
    p_t1 = tf_t1.paragraphs[0]
    p_t1.text = "🏏 Real-Time Concurrent IPL Auction System"
    p_t1.font.size = Pt(34)
    p_t1.font.bold = True
    p_t1.font.color.rgb = TEXT_LIGHT
    p_t1.font.name = FONT_TITLE

    p_sub = tf_t1.add_paragraph()
    p_sub.text = "A Multi-Franchise Digital Bidding Arena with Guaranteed Concurrency & Financial Integrity"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = CYAN
    p_sub.font.name = FONT_BODY

    members = [
        ("Srijan Srivastava", "1240440037", "Project Leader", "Core Auction Engine"),
        ("Amit Kumar Rajput", "1240440008", "Security Specialist", "Identity & Auth"),
        ("Anshika", "1240440012", "Treasury Specialist", "Budgets & Rules"),
        ("Akhilesh Kumar Pal", "1240440006", "UI/UX Designer", "War Room Cockpit"),
        ("Suryansh Sahu", "1240440039", "Live Sync Lead", "WebSockets & QA")
    ]

    card_width = Inches(2.15)
    gap = Inches(0.14)
    start_left = Inches(1.0)
    top_pos = Inches(3.6)

    for i, (name, roll, role, duty) in enumerate(members):
        col_left = start_left + i * (card_width + gap)
        box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col_left, top_pos, card_width, Inches(2.8))
        box.fill.solid()
        box.fill.fore_color.rgb = PANEL_BG
        box.line.color.rgb = CYAN if i == 0 else PANEL_BORDER
        box.line.width = Pt(1.5 if i == 0 else 1.0)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.15)
        tf.margin_top = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = f"MEMBER {i+1}"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = GOLD
        p.font.name = FONT_TITLE
        p.alignment = PP_ALIGN.CENTER

        p_name = tf.add_paragraph()
        p_name.text = name
        p_name.font.size = Pt(13)
        p_name.font.bold = True
        p_name.font.color.rgb = TEXT_LIGHT
        p_name.font.name = FONT_TITLE
        p_name.alignment = PP_ALIGN.CENTER
        p_name.space_after = Pt(4)

        p_roll = tf.add_paragraph()
        p_roll.text = f"Roll: {roll}"
        p_roll.font.size = Pt(9.5)
        p_roll.font.color.rgb = TEXT_MUTED
        p_roll.font.name = FONT_BODY
        p_roll.alignment = PP_ALIGN.CENTER
        p_roll.space_after = Pt(8)

        p_role = tf.add_paragraph()
        p_role.text = role
        p_role.font.size = Pt(11)
        p_role.font.bold = True
        p_role.font.color.rgb = CYAN
        p_role.alignment = PP_ALIGN.CENTER

        p_duty = tf.add_paragraph()
        p_duty.text = duty
        p_duty.font.size = Pt(10)
        p_duty.font.color.rgb = TEXT_LIGHT
        p_duty.alignment = PP_ALIGN.CENTER

    # SLIDE 2: THE PROBLEM (Srijan)
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The Real-World Dilemma: When Crores Collide", "Problem & Solution", "Srijan", 2)

    create_card(s2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "The Auction Room Challenge",
        [
            "⚡ Split-Second Clashes: Multiple teams raise paddles at the exact same millisecond.",
            "💸 Huge Stakes: Franchise purses exceed ₹100 Crore. A single lag ruins an entire season.",
            "📡 10 Distributed Screens: All team tables must see the exact same price at all times."
        ],
        badge_text="REAL WORLD"
    )

    create_card(s2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "What Fails in Ordinary Apps",
        [
            "❌ Colliding Bids: Two teams awarded the player at the exact same valuation.",
            "❌ Wallet Desync: A team's budget updates on one laptop but stays stale on the big screen.",
            "❌ The 'F5 Trap': Having to refresh the webpage causes owners to miss the 15-second window."
        ],
        badge_text="THE RISK",
        border_color=RED
    )

    create_card(s2, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
        "Our Solution: Digital Referee",
        [
            "🎯 Single-Queue Referee: Concurrent bids are lined up in microsecond order. Only 1 bid wins.",
            "🔒 Mathematical Safety: Purses update instantly with zero math or currency errors.",
            "⚡ < 50ms Broadcast: All 10 screens update in real time with zero manual refreshing."
        ],
        badge_text="THE FIX",
        border_color=GREEN
    )

    # SLIDE 3: SYSTEM ARCHITECTURE (Srijan)
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "How the Auction Hall Works Behind the Scenes", "System Architecture", "Srijan", 3)

    arch_layers = [
        ("Layer 1: War Rooms", "What Owners See", [
            "• Cockpit themed with franchise colors",
            "• 1-Click quick paddles (+₹20L, +₹50L, +₹1Cr)",
            "• Live purse fuel gauge and active valuation"
        ], CYAN),
        ("Layer 2: Digital Referee", "The Core Engine", [
            "• Inspects every incoming bid in microseconds",
            "• Validates budget caps and auction rules",
            "• Resolves race conditions between competing bids"
        ], GOLD),
        ("Layer 3: Broadcast Tower", "Instant Sync", [
            "• Pushes updates to all 10 laptops in < 50ms",
            "• Broadcasts live countdown timer to all screens",
            "• Powers the real-time Orbit Arena visuals"
        ], GREEN),
        ("Layer 4: Permanent Vault", "Audit & Security", [
            "• Encrypted player contracts and team purse ledgers",
            "• Permanent audit trail for every hammer-down",
            "• Zero possibility of post-auction tampering"
        ], CYAN)
    ]

    card_w3 = Inches(2.7)
    gap3 = Inches(0.2)
    left3 = Inches(0.8)

    for i, (title, sub, bullets, color) in enumerate(arch_layers):
        c_left = left3 + i * (card_w3 + gap3)
        create_card(s3, c_left, Inches(1.8), card_w3, Inches(5.0),
            title,
            [f"Role: {sub}"] + bullets,
            badge_text=f"LAYER {i+1}",
            border_color=color
        )

    # SLIDE 4: SECURITY & ACCESS CONTROL (Amit)
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Digital Passports: Preventing Impersonation & Cheating", "Security & Fair Play", "Amit", 4)

    create_card(s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Tamper-Proof Digital Passports",
        [
            "🔑 Encrypted Login Badges: Every franchise receives a cryptographic badge on login.",
            "🛡️ Verified Paddle Ownership: Before a bid registers, the system verifies the team's identity.",
            "🚫 Zero Cross-Team Bidding: Chennai cannot bid with Mumbai's money; Kolkata cannot touch Bangalore's paddle.",
            "🔒 Protected Arena: Unauthorized spectators cannot join or disturb the live auction room."
        ],
        badge_text="IDENTITY SHIELD"
    )

    create_card(s4, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0),
        "Strict Role Separation",
        [
            "⚖️ BCCI Auctioneer Desk (Admin):",
            "  • Introduces new cricketers to the center stage",
            "  • Controls the countdown clock (Hold, Start, Pause)",
            "  • Declares 'Sold' with Hammer Down or marks 'Unsold'",
            "",
            "🏏 Franchise War Rooms (Owners):",
            "  • Pure bidding interface restricted strictly to own team funds",
            "  • Complete view of live rival actions and purse levels"
        ],
        badge_text="TWO TIERS",
        border_color=GOLD
    )

    # SLIDE 5: AUCTION RULES & PURSE SAFEGUARDS (Anshika)
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "The Rulebook: Mathematical Guardrails & Budget Caps", "Auction Rules & Integrity", "Anshika", 5)

    rules = [
        ("Rule 1: Hard Purse Limit", [
            "• Every team begins with a strict purse (e.g. ₹100 Crore).",
            "• If remaining purse is ₹12 Cr, a bid of ₹13 Cr is physically blocked.",
            "• Zero overdraft risk — budgets are protected automatically."
        ]),
        ("Rule 2: Anti-Self Bidding", [
            "• A franchise holding the top bid cannot bid against itself.",
            "• Stops accidental double-clicks or artificial price inflation.",
            "• Another franchise must raise a paddle before you can bid again."
        ]),
        ("Rule 3: Price Progression", [
            "• Opening bid must meet or exceed the player's base price.",
            "• Each subsequent bid must strictly increase the active valuation.",
            "• Standardized increment brackets eliminate confusion."
        ]),
        ("Rule 4: Irreversible Gavel", [
            "• Once the auctioneer hits 'Hammer Down', the sale is locked.",
            "• Funds are permanently deducted from the buying team's purse.",
            "• Player roster assignment is final and auditable."
        ])
    ]

    for i, (rtitle, ritems) in enumerate(rules):
        row = i // 2
        col = i % 2
        r_left = Inches(0.8) + col * Inches(6.0)
        r_top = Inches(1.8) + row * Inches(2.5)
        create_card(s5, r_left, r_top, Inches(5.7), Inches(2.3),
            rtitle, ritems, badge_text=f"RULE #{i+1}", border_color=CYAN
        )

    # SLIDE 6: PLAYER LIFECYCLE (Anshika)
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "The Player Lifecycle: From Draft to Final Hammer", "Auction State Machine", "Anshika", 6)

    stages = [
        ("STAGE 1: WAITING (ON HOLD)", [
            "• Cricketer is brought to the stage for introduction.",
            "• Timer is still and locked at 15.00s.",
            "• Bidding paddles are disabled to prevent premature clicks."
        ], GOLD),
        ("STAGE 2: LIVE BIDDING", [
            "• Auctioneer clicks 'START TIMER' to begin the lot.",
            "• 15-second clock turns green and counts down.",
            "• Franchises actively bid; each bid resets clock to 15s."
        ], GREEN),
        ("STAGE 3: SOLD / UNSOLD", [
            "• SOLD: Auctioneer hammers down ➔ Team purse deducted.",
            "• UNSOLD: Clock expires with no bids ➔ Returned to pool.",
            "• Lot is finalized; system advances cleanly to next lot."
        ], CYAN)
    ]

    for i, (stitle, sitems, scolor) in enumerate(stages):
        s_left = Inches(0.8) + i * Inches(3.95)
        create_card(s6, s_left, Inches(1.8), Inches(3.8), Inches(5.0),
            stitle, sitems, badge_text=f"PHASE {i+1}", border_color=scolor
        )

    # SLIDE 7: WAR ROOM DASHBOARD (Akhilesh & Suryansh)
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "The Franchise Cockpit & Zero-Refresh Broadcasting", "War Room Experience & Sync", "Akhilesh & Suryansh", 7)

    create_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Franchise Cockpit (Akhilesh)",
        [
            "🎨 Franchise-Themed Immersion:",
            "  • CSK turns yellow with lion badge, MI turns blue, KKR purple.",
            "",
            "⚡ 1-Click Tactical Bidding Paddles:",
            "  • +₹20 Lakhs | +₹50 Lakhs | +₹1 Crore instant clicks.",
            "  • Eliminates typing errors during high-pressure 15s rounds.",
            "",
            "🔋 Purse Containment Fuel Gauge:",
            "  • Shows live remaining budget in Crores at all times.",
            "",
            "🛰️ Orbit Arena Visualization:",
            "  • 10 team logos orbit the player; active bidder pulses live."
        ],
        badge_text="UI / UX LEAD"
    )

    create_card(s7, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0),
        "Live Push Sync & Timer (Suryansh)",
        [
            "⚡ Zero-Refresh Technology:",
            "  • No owner ever touches 'F5' or refresh buttons.",
            "  • Price updates push to all 10 laptops in < 50 milliseconds.",
            "",
            "⏱️ The Controlled Auctioneer Timer:",
            "  • Player lots load on 'ON HOLD' (stays still at 15.00s).",
            "  • Auctioneer introduces player, then clicks 'START TIMER'.",
            "  • Ticking is synchronized live across all screens simultaneously.",
            "  • Every new bid resets timer back to 15s for counter-bids."
        ],
        badge_text="LIVE SYNC LEAD",
        border_color=GREEN
    )

    # SLIDE 8: LIVE DEMO & ROADMAP (Suryansh)
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Live Demonstration & System Walkthrough", "Live Demo & Summary", "Suryansh", 8)

    create_card(s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Live Demonstration Steps",
        [
            "1️⃣ Host Setup: Open Screen 1 (Auctioneer) & Screen 2 (CSK Owner).",
            "2️⃣ Still Timer: Point out timer is frozen at 15.00s on 'ON HOLD'.",
            "3️⃣ Start Lot: Auctioneer clicks 'START TIMER' ➔ Clock turns green on both laptops.",
            "4️⃣ Raise Bid: CSK clicks '+₹20 Lakhs' ➔ Valuation jumps and timer resets.",
            "5️⃣ Final Call: Auctioneer calls 'Going once, twice...' ➔ Clicks 'HAMMER DOWN'!",
            "6️⃣ Celebration: Player marked SOLD to CSK, confetti fires, purse deducts."
        ],
        badge_text="DEMO PLAN",
        border_color=GOLD
    )

    create_card(s8, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0),
        "Key Project Takeaways",
        [
            "🏆 Zero Concurrency Conflicts: Microsecond resolution ensures bids never collide.",
            "🛡️ Total Financial Integrity: Hard purse caps make budget overdrafts impossible.",
            "⚡ Millisecond-Grade Speed: < 50ms updates deliver a real-world broadcast experience.",
            "🎯 True Multi-Franchise War Room: Built for high-pressure decisions with 1-click controls.",
            "",
            "✨ Thank you, judges! We now welcome your questions."
        ],
        badge_text="SUMMARY",
        border_color=CYAN
    )

    out_path = r"c:\Users\hp\ipl_auction_system complete\IPL-Auction-NonTech-Presentation.pptx"
    prs.save(out_path)
    print(f"Successfully generated non-tech PPT at: {out_path}")

if __name__ == "__main__":
    create_deck()
